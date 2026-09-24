from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BG=RGBColor(0x0B,0x12,0x20); PANEL=RGBColor(0x12,0x1C,0x30); LINE=RGBColor(0x1F,0x3A,0x5F)
BLUE=RGBColor(0x3B,0x82,0xF6); CYAN=RGBColor(0x22,0xD3,0xEE); TEXT=RGBColor(0xE5,0xED,0xF7)
DIM=RGBColor(0x94,0xA3,0xB8); GREEN=RGBColor(0x34,0xD3,0x99); AMBER=RGBColor(0xFB,0xBF,0x24); RED=RGBColor(0xF8,0x71,0x71)
FONT='PingFang SC'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]
total=[0]

def txt(slide,x,y,w,h,text,size=16,color=TEXT,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=Inches(0.05)
    for i,line in enumerate(text if isinstance(text,list) else [text]):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        r=p.add_run(); r.text=line; r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold; r.font.name=FONT
    return tb

def box(slide,x,y,w,h,fill=PANEL,line=LINE,shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s=slide.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h)); s.fill.solid(); s.fill.fore_color.rgb=fill
    s.line.color.rgb=line; s.line.width=Pt(1)
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0]=0.08
    s.shadow.inherit=False; return s

def slide(title,kicker=None):
    s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
    bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(0.6),Inches(0.55),Inches(0.08),Inches(0.62)); bar.fill.solid(); bar.fill.fore_color.rgb=CYAN; bar.line.fill.background()
    if kicker: txt(s,0.8,0.35,10,0.35,kicker,12,CYAN,True)
    txt(s,0.8,0.62,11.5,0.7,title,28,TEXT,True)
    total[0]+=1
    txt(s,11.8,7.0,1.2,0.3,f'{total[0]:02d}',11,DIM,align=PP_ALIGN.RIGHT)
    txt(s,0.6,7.0,6,0.3,'Aperture · AI Native SDLC Control Plane',11,DIM)
    return s

def card(s,x,y,w,h,head,lines,accent=BLUE,size=14):
    box(s,x,y,w,h)
    a=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y+0.18),Inches(0.06),Inches(0.42)); a.fill.solid(); a.fill.fore_color.rgb=accent; a.line.fill.background()
    txt(s,x+0.22,y+0.15,w-0.4,0.5,head,17,TEXT,True)
    txt(s,x+0.22,y+0.7,w-0.4,h-0.8,['· '+l for l in lines],size,DIM)

def bullets(s,x,y,w,items,size=17,gap=0.62):
    for i,(head,body) in enumerate(items):
        yy=y+i*gap
        d=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(yy+0.12),Inches(0.14),Inches(0.14)); d.fill.solid(); d.fill.fore_color.rgb=CYAN; d.line.fill.background()
        tb=txt(s,x+0.3,yy,w,gap,'',size)
        p=tb.text_frame.paragraphs[0]
        r=p.runs[0]; r.text=head; r.font.bold=True; r.font.color.rgb=TEXT
        r2=p.add_run(); r2.text='  '+body; r2.font.size=Pt(size-2); r2.font.color.rgb=DIM; r2.font.name=FONT

def flow(s,y,steps,x0=0.6,width=12.1,h=1.1,colors=None):
    n=len(steps); gap=0.28; w=(width-gap*(n-1))/n
    for i,(head,sub) in enumerate(steps):
        x=x0+i*(w+gap); c=(colors[i] if colors else BLUE)
        b=box(s,x,y,w,h,PANEL,c)
        txt(s,x+0.08,y+0.12,w-0.16,0.4,head,15,TEXT,True,PP_ALIGN.CENTER)
        txt(s,x+0.08,y+0.52,w-0.16,h-0.55,sub,11,DIM,align=PP_ALIGN.CENTER)
        if i<n-1:
            ar=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x+w+0.03),Inches(y+h/2-0.1),Inches(0.22),Inches(0.2)); ar.fill.solid(); ar.fill.fore_color.rgb=CYAN; ar.line.fill.background()

def notes(s,text): s.notes_slide.notes_text_frame.text=text

# 1 Title
s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG; total[0]+=1
for i,c in enumerate([BLUE,CYAN,LINE]):
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(0.8+i*0.5),Inches(1.9),Inches(0.35),Inches(0.08)); r.fill.solid(); r.fill.fore_color.rgb=c; r.line.fill.background()
txt(s,0.8,2.2,11.5,1.0,'Aperture 控制平面',48,TEXT,True)
txt(s,0.8,3.25,11.5,0.7,'让小团队安全、快速地接受 AI Agent 产出的代码变更',24,CYAN)
txt(s,0.8,4.2,11.5,0.5,'演示说明 · 面向 3–8 人研发团队 · 自托管 · Agent 中立',16,DIM)
txt(s,0.8,6.6,11.5,0.4,'2026-09 · 研究原型 v0.1',12,DIM)
notes(s,'开场：一句话说明平台定位——不是又一个写代码的 Agent，而是管住 Agent 产出、让人能快速放心地接受变更的控制层。')

# 2 Problem
s=slide('AI 写代码很快，人审查跟不上','为什么需要它')
card(s,0.6,1.6,5.9,2.5,'首要缺口：审查带宽',['Agent 一小时产出的变更，审查者要读很久','证据散落在终端、CI、聊天记录里','"它说测试通过了" ≠ 测试真的独立通过','结果：要么橡皮章审批，要么审查成为瓶颈'],RED,15)
card(s,6.8,1.6,5.9,2.5,'随之而来的四个缺口',['意图：模糊 Prompt 直接交给 Agent','上下文：Agent 用了哪些文件、规则，无从复现','控制：权限、网络、凭据约束不透明','评估与证据：验收标准和检查没有对上'],AMBER,15)
box(s,0.6,4.45,12.1,1.9,PANEL,BLUE)
txt(s,0.9,4.6,11.5,0.5,'北极星：让团队里任何一个人，都能在几分钟内安全地接受另一个人的 Agent 产出的变更',19,TEXT,True)
txt(s,0.9,5.25,11.5,1.0,['第一阶段唯一必达目标：降低单次变更的审查成本，并让审查负载在团队内可见。','默认 AI 执行、人审查合并；低风险且证据充分的变更，后续逐步开放自动合并试验。'],15,DIM)
notes(s,'强调问题是审查带宽，而不是生成能力。')

# 3 Loop
s=slide('一条从意图到合并的受控闭环','核心流程')
flow(s,1.55,[('Intent','目标 · 约束 · 验收标准 · 风险'),('批准 Intent','中高风险需非起草人批准'),('Agent Run','在隔离 worktree 中执行'),('检查与评估','结果绑定 Head SHA')],h=0.95,colors=[BLUE,BLUE,CYAN,CYAN])
flow(s,2.7,[('Evidence','可复验的证据包'),('非作者审查','逐条验收标准门禁'),('合并证据','目标分支 = 批准的 Head'),('发布候选','独立发布批准')],h=0.95,colors=[CYAN,GREEN,GREEN,GREEN])
card(s,0.6,3.95,3.9,2.4,'人负责决定',['起草并批准意图','审查证据、批准或请求修改','合并、发布、推翻门禁'],BLUE,15)
card(s,4.7,3.95,3.9,2.4,'Agent 负责执行',['在隔离工作区改代码','产出分支与提交','不持有托管凭据'],CYAN,15)
card(s,8.8,3.95,3.9,2.4,'平台负责记录与强制',['每一步写入哈希链事件','规则在服务端执行','证据与决策绑定到具体 SHA'],GREEN,15)
notes(s,'带观众走一遍闭环；后面每一页演示对应其中一段。')

# 4 Principles
s=slide('设计原则','这些原则决定了产品取舍')
bullets(s,0.8,1.7,11.5,[
 ('证据先于信任','审批必须绑定具体的 Head SHA、检查 ID 与证据摘要；换了 Head，旧审批自动失效'),
 ('在边界强制','所有规则在服务端数据库层执行，界面只是投影，脚本和 HTTP 都绕不过去'),
 ('安全的默认值','验收标准不写标注默认按「关键 · 确定性」处理，降级必须显式声明'),
 ('分级自治','低风险走轻量路径，高风险需要人工关键标准；不搞"所有变更一律人审"'),
 ('附着而非替代','贴在现有 Git / GitHub 上工作，不另建一套代码托管或需求系统'),
 ('自托管、Agent 中立','Codex、其他 CLI Agent 或容器都可以接入；不依赖外部 SaaS'),
 ('失败如实保留','失败的 Run 保留在事件日志里，不删除、不伪装成功'),
],17,0.72)

# 5 Demo prep
s=slide('演示准备','开始之前')
card(s,0.6,1.6,6.3,2.9,'启动服务',['cd control-plane/workbench','npm install && npm run build','node --experimental-strip-types server/main.ts','浏览器打开 http://127.0.0.1:8787'],BLUE,15)
card(s,7.1,1.6,5.6,2.9,'准备三个角色',['Owner：创建项目、管理成员、合并','Developer：起草 Intent、启动 Run','Reviewer：审查证据、批准','建议用两个浏览器窗口分别登录作者与审查人'],CYAN,15)
card(s,0.6,4.75,12.1,1.65,'可选：接入 Codex Builder 与 GitHub',['设置 CONTROL_PLANE_AGENT_EXECUTABLE / ARGS_JSON 指向 codex-builder.mjs，Run 由真实 Codex 执行','GitHub 项目：启动前 export APERTURE_GITHUB_TOKEN=…，项目里只填变量名'],GREEN,14)
notes(s,'演示时作者和审查人必须是不同账号，否则会被"不能批准自己"的规则拦下，这本身也可以作为演示点。')

# 6 Projects
s=slide('第 1 步 · 创建项目，选择代码托管','演示')
card(s,0.6,1.6,3.9,3.2,'多项目',['侧栏切换项目，数据按项目隔离','非成员请求其他项目返回 404','一个项目对应一个仓库'],BLUE,15)
card(s,4.7,1.6,3.9,3.2,'按项目授予角色',['Maintainer / Reviewer / Developer','同一人在不同项目可有不同角色','角色变更是带身份的决策事件'],CYAN,15)
card(s,8.8,1.6,3.9,3.2,'代码托管可选',['本地 Git：平台合并','GitHub（含 Enterprise）：推分支、开 PR、回写门禁','Token 只存环境变量名'],GREEN,15)
box(s,0.6,5.05,12.1,1.3,PANEL,AMBER)
txt(s,0.9,5.2,11.6,1.1,['演示点：在「项目」页新建项目 → 测试连接 → 把成员加入项目并分配角色 → 用成员账号登录，切换器只显示他所在的项目。','Run 与提案表单不再有仓库路径输入，仓库只从项目解析。'],15,TEXT)

# 7 Intent
s=slide('第 2 步 · 起草并批准 Intent','演示')
card(s,0.6,1.6,6.2,3.0,'结构化意图',['目标、约束、风险等级、验收标准','「套用模版」按产品类型填充，占位符必须逐项填完','实时解析预览 + 模糊度提示（不阻塞）','内容进入 contentDigest 与 Agent prompt'],BLUE,15)
box(s,7.0,1.6,5.7,3.0)
txt(s,7.25,1.75,5.3,0.4,'验收标准标注',17,TEXT,True)
rows=[('[确定性]','测试 / 退出码证明（默认）'),('[模型]','由模型评估证明'),('[人工]','由审查人判断并写意见'),('[关键]','未证明不得批准（默认）'),('[参考]','非关键，不阻塞合并')]
for i,(k,v) in enumerate(rows):
    txt(s,7.25,2.3+i*0.44,1.4,0.4,k,14,CYAN,True); txt(s,8.7,2.3+i*0.44,3.9,0.4,v,14,DIM)
card(s,0.6,4.85,12.1,1.55,'服务端硬规则',['高风险必须有一条关键的 [人工] 标准；中高风险 Intent 需非起草人批准；新版本使旧批准失效；未批准的 Intent 无法启动 Run'],RED,14)

# 8 Run
s=slide('第 3 步 · 启动 Agent Run','演示')
flow(s,1.7,[('准入','Intent 已批准 · 项目仓库'),('隔离工作区','git worktree · 环境白名单'),('Codex 执行','按验收标准改代码'),('独立检查','node-tests · @baseline 重跑'),('证据包','stdout/stderr 摘要 · 覆盖映射')],h=1.3,colors=[BLUE,CYAN,CYAN,GREEN,GREEN])
card(s,0.6,3.35,6.0,3.0,'防止"自己给自己打分"',['只靠 Agent 自己写的测试通过 = 仅自带测试','对关键标准，它和证据失败一样阻塞批准','解除方式：@baseline 重置测试文件后重跑，或声明 testPaths 且 Agent 未改动'],AMBER,14)
card(s,6.8,3.35,5.9,3.0,'运行中心能看到',['排队、运行、取消，实时事件流','Runtime Attestation：隔离级别、环境变量名','策略文件 .aperture/ 被改动会被标记','失败 Run 与原因完整保留'],BLUE,14)

# 9 Review
s=slide('第 4 步 · 审查：证据先展开，再决策','演示')
card(s,0.6,1.6,3.9,2.4,'审查人分配',['按负载自动分配或手动指定','作者永远不能被分配','有分配时只有被分配人能决策'],BLUE,14)
card(s,4.7,1.6,3.9,2.4,'逐条证据门禁',['每条关键标准都要有对应的通过证据','待证据 / 失败 / 仅自带测试 → 不能批准','[人工] 标准必须写审查意见'],CYAN,14)
card(s,8.8,1.6,3.9,2.4,'治理决策',['Override：只推翻一条标准，只对当前 Head','Reject：需要理由，提案关闭','完整性检查永远不能推翻'],RED,14)
box(s,0.6,4.25,12.1,2.1,PANEL,GREEN)
txt(s,0.9,4.4,11.6,0.4,'演示点',16,GREEN,True)
txt(s,0.9,4.85,11.6,1.5,['1. 审查人打开提案，先展开证据包（平台会重新读取文件并校验摘要）','2. 查看验收标准覆盖：哪条由哪个检查证明、状态是什么','3. 批准 → 审批绑定 Head SHA、检查 ID、证据摘要；作者再推新提交，审批自动失效'],15,TEXT)

# 10 Merge
s=slide('第 5 步 · 合并、同步 GitHub 与发布','演示')
card(s,0.6,1.6,4.0,3.0,'合并证据',['目标分支 SHA 必须等于批准的 Head','复核检查、证据、审批仍然有效','Merge Evidence 追加写入，可复验摘要'],BLUE,14)
card(s,4.8,1.6,4.0,3.0,'GitHub 项目',['平台推 aperture/* 分支并开 PR','aperture/gate 状态随审批与证据变化','GitHub Actions 结果导入为独立检查'],CYAN,14)
card(s,9.0,1.6,3.7,3.0,'发布候选',['从合并证据创建源码快照','需要另一位 Owner / Maintainer 批准发布'],GREEN,14)
box(s,0.6,4.85,5.95,1.5,PANEL,BLUE)
txt(s,0.85,4.95,5.6,0.4,'合并方式 · control_plane',15,TEXT,True)
txt(s,0.85,5.4,5.6,1.0,'平台合并后以 force-with-lease 推送；远程已前进或被分支保护拒绝时回滚，提案保持已批准',13,DIM)
box(s,6.75,4.85,5.95,1.5,PANEL,AMBER)
txt(s,7.0,4.95,5.6,0.4,'合并方式 · host_protected',15,TEXT,True)
txt(s,7.0,5.4,5.6,1.0,'在 GitHub 上合并，平台同步并写入合并证据；绕过门禁的合并标红并留下审计事件',13,DIM)

# 11 Real case
s=slide('真实案例：Import Validator','已在本机用真实 Codex 跑通')
card(s,0.6,1.6,4.0,2.6,'任务',['只允许 CSV / JSON','最大 5 MiB，拒绝负数与非整数','错误顺序确定，不削弱既有测试'],BLUE,14)
box(s,4.8,1.6,3.9,2.6)
txt(s,5.0,1.75,3.6,0.4,'结果',17,TEXT,True)
txt(s,5.0,2.3,3.6,0.8,'1/4 → 6/6',40,GREEN,True)
txt(s,5.0,3.2,3.6,0.9,'基线 1 项通过、3 项失败；Codex 修改 2 个文件后 6 项全部通过',13,DIM)
card(s,8.9,1.6,3.8,2.6,'治理链',['created → check → evidence','→ viewed → approved','两条事件链完整性复验通过'],GREEN,14)
card(s,0.6,4.45,12.1,1.9,'过程中暴露并修复的两个问题（失败 Run 均保留）',['Codex CLI 不允许同时使用 --sandbox workspace-write 与 --approve-for-me → Wrapper 调整参数','Process Runtime 只传 PATH/HOME，丢失 CODEX_* 环境 → 改为显式白名单继承，证明里只记录变量名'],AMBER,14)

# 12 Trust
s=slide('可信与安全','审计者关心的问题')
bullets(s,0.8,1.7,11.5,[
 ('哈希链事件账本','每个聚合的事件都带前一事件摘要，篡改可被检测；事件表只追加'),
 ('决策身份','团队模式下决策需要 GitHub 登录的已验证身份（OAuth + PKCE），身份冻结进事件'),
 ('凭据不落盘','托管 Token 只存变量名；不进数据库、argv、git 配置、日志或浏览器；测试会扫描数据目录'),
 ('Worker 不持有凭据','推送、开 PR、同步都在主进程完成，Agent 进程拿不到 Token'),
 ('系统身份分离','GitHub 报告的检查、合并、关闭以 system:code-host 身份记录，与人的决策区分'),
 ('迁移可验证','历史库升级后事件链、合并证据、发布摘要全部复验'),
],17,0.8)

# 13 Status
s=slide('当前状态与已知限制','如实说明')
card(s,0.6,1.6,6.0,4.75,'已完成并验证',['npm run check 全绿：30+ 条端到端 smoke','多项目、按项目角色、本地 Git / GitHub 托管','Intent 模版与验收标准证据门禁','审查分配、Override / Reject、策略文件标记','外部身份（GitHub 登录）与团队模式','LLM Provider 配置、Codex Builder 真实案例'],GREEN,15)
card(s,6.8,1.6,5.9,4.75,'已知限制',['服务只监听 127.0.0.1，尚不能多人远程使用','真实 GitHub 端到端测试脚本已就绪，待运行','Run 结束后 worktree 尚未自动清理','没有通知、Webhook、GitLab','验收标准到检查的显式映射（verifiedBy）未实现'],AMBER,15)

# 14 Next
s=slide('下一步','路线')
flow(s,1.8,[('真机验证','真实 GitHub e2e · worktree 清理'),('团队可用','内网部署 · HTTPS · 备份恢复'),('通知','审查分配、门禁失败、绕过门禁'),('治理补齐','verifiedBy · 按项目信任模式'),('按需扩展','GitLab · Webhook · 自动合并试验')],h=1.5,colors=[BLUE,CYAN,CYAN,GREEN,GREEN])
box(s,0.6,4.0,12.1,2.3,PANEL,BLUE)
txt(s,0.9,4.2,11.5,0.5,'需要团队决定',18,TEXT,True)
txt(s,0.9,4.8,11.5,1.5,['· 访问方式：内网、公网还是 VPN —— 决定部署方案','· 通知渠道：飞书、Slack、企业微信还是邮件','· 首批试用项目与审查人名单'],16,DIM)

prs.save('/Users/wangzhen/Documents/AI/AI Native SDLC/control-plane/docs/Aperture_平台演示说明.pptx')
print('slides', len(prs.slides))
